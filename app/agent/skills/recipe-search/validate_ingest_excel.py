#!/usr/bin/env python3
"""
入库前 Excel 数据校验/清洗。

默认只校验并打印不合规行；传 --out 时会保存删除不合规行后的副本。

用法：
  .venv/bin/python validate_ingest_excel.py tmp/recipes.xlsx
  .venv/bin/python validate_ingest_excel.py tmp/recipes.xlsx --out tmp/recipes.validated.xlsx
  .venv/bin/python validate_ingest_excel.py tmp/recipes.xlsx --out tmp/recipes.validated.xlsx --report tmp/invalid_rows.csv

默认必填且非空：
  id / 名称 / 食材 / 图片地址 / lang

唯一性规则：
  业务主键暂定为 id + lang。
  同一个 id 可以同时存在 zh/en 两行；只有同一个业务主键重复才算不合规。
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import openpyxl


COLUMN_MAP = {
    "recipe_id": "id",
    "name": "名称",
    "ingredients": "食材",
    "tags": "标签",
    "image_url": "图片地址",
    "lang": "lang",
}

REQUIRED_FIELDS = ("recipe_id", "name", "ingredients", "image_url", "lang")
BUSINESS_KEY_FIELDS = ("recipe_id", "lang")
VALID_LANGS = {"zh", "en"}
SPLIT_RE = re.compile(r"[、,，;；/\|\r\n]+")
BLANK_STRINGS = {"", "nan", "none", "null", "未知", "未知食材", "n/a", "na"}


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in BLANK_STRINGS
    return False


def _split_multi(value: Any) -> list[str]:
    if _is_blank(value):
        return []
    if isinstance(value, (list, tuple, set)):
        return [str(x).strip() for x in value if not _is_blank(x)]
    return [p.strip() for p in SPLIT_RE.split(str(value)) if p.strip() and p.strip().lower() not in BLANK_STRINGS]


def _valid_url(value: Any) -> bool:
    if _is_blank(value):
        return False
    parsed = urlparse(str(value).strip())
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def _headers(ws) -> dict[str, int]:
    values = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    return {name: idx + 1 for idx, name in enumerate(values) if name}


def _cell(row, col_idx: int):
    return row[col_idx - 1].value if col_idx and col_idx - 1 < len(row) else None


def _row_values(row, header_idx: dict[str, int]) -> dict[str, Any]:
    return {logical: _cell(row, header_idx.get(col_name, 0)) for logical, col_name in COLUMN_MAP.items()}


def _business_key(values: dict[str, Any]) -> tuple[str, str] | None:
    recipe_id = "" if _is_blank(values.get("recipe_id")) else str(values.get("recipe_id")).strip()
    lang = "" if _is_blank(values.get("lang")) else str(values.get("lang")).strip().lower()
    if not recipe_id or lang not in VALID_LANGS:
        return None
    return recipe_id, lang


def _business_key_text(values: dict[str, Any]) -> str:
    key = _business_key(values)
    return "" if key is None else f"{key[0]}:{key[1]}"


def _validate_row(values: dict[str, Any], seen_business_keys: set[tuple[str, str]], require_tags: bool) -> list[str]:
    reasons: list[str] = []

    for field in REQUIRED_FIELDS:
        if field == "ingredients":
            if not _split_multi(values.get(field)):
                reasons.append("missing_ingredients")
        elif _is_blank(values.get(field)):
            reasons.append(f"missing_{field}")

    lang = "" if _is_blank(values.get("lang")) else str(values.get("lang")).strip().lower()
    if lang and lang not in VALID_LANGS:
        reasons.append("invalid_lang")

    business_key = _business_key(values)
    if business_key:
        if business_key in seen_business_keys:
            reasons.append("duplicate_business_key")
        seen_business_keys.add(business_key)

    if not _is_blank(values.get("image_url")) and not _valid_url(values.get("image_url")):
        reasons.append("invalid_image_url")

    if require_tags and not _split_multi(values.get("tags")):
        reasons.append("missing_tags")

    return reasons


def _write_report(path: Path, invalid_rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["excel_row", "business_key", "id", "name", "lang", "reasons", "ingredients", "image_url"],
        )
        writer.writeheader()
        for item in invalid_rows:
            values = item["values"]
            writer.writerow({
                "excel_row": item["excel_row"],
                "business_key": _business_key_text(values),
                "id": values.get("recipe_id"),
                "name": values.get("name"),
                "lang": values.get("lang"),
                "reasons": "|".join(item["reasons"]),
                "ingredients": values.get("ingredients"),
                "image_url": values.get("image_url"),
            })


def main() -> int:
    parser = argparse.ArgumentParser(description="校验入库 Excel，并可输出删除不合规行后的副本")
    parser.add_argument("xlsx", help="待校验 Excel")
    parser.add_argument("--sheet", help="工作表名；默认 active sheet")
    parser.add_argument("--out", help="输出清洗后的 Excel 副本；会删除不合规行")
    parser.add_argument("--in-place", action="store_true", help="直接修改源文件；会先生成 .bak 备份")
    parser.add_argument("--report", help="输出不合规行 CSV 报告")
    parser.add_argument("--sample", type=int, default=20, help="最多打印多少条样例")
    parser.add_argument("--require-tags", action="store_true", help="把 标签 列也作为必填")
    args = parser.parse_args()

    src = Path(args.xlsx)
    if not src.exists():
        print(f"❌ 找不到 Excel：{src}")
        return 2
    if args.out and args.in_place:
        print("❌ --out 和 --in-place 不能同时使用")
        return 2

    wb = openpyxl.load_workbook(src)
    ws = wb[args.sheet] if args.sheet else wb.active
    header_idx = _headers(ws)

    missing_columns = [col for col in COLUMN_MAP.values() if col not in header_idx]
    if missing_columns:
        print(f"❌ Excel 缺少列：{missing_columns}")
        print(f"   当前列：{list(header_idx.keys())}")
        return 2

    seen_business_keys: set[tuple[str, str]] = set()
    invalid_rows: list[dict] = []
    valid_count = 0

    for row in ws.iter_rows(min_row=2):
        if all(c.value is None for c in row):
            continue
        values = _row_values(row, header_idx)
        reasons = _validate_row(values, seen_business_keys, args.require_tags)
        if reasons:
            invalid_rows.append({
                "excel_row": row[0].row,
                "values": values,
                "reasons": reasons,
            })
        else:
            valid_count += 1

    total = valid_count + len(invalid_rows)
    reason_counts: dict[str, int] = {}
    for item in invalid_rows:
        for reason in item["reasons"]:
            reason_counts[reason] = reason_counts.get(reason, 0) + 1

    print(f"Excel: {src}")
    print(f"Sheet: {ws.title}")
    print(f"Business key: {' + '.join(BUSINESS_KEY_FIELDS)}")
    print(f"Rows: total={total}, valid={valid_count}, invalid={len(invalid_rows)}")
    print(f"Reasons: {reason_counts}")
    for item in invalid_rows[:args.sample]:
        values = item["values"]
        print("  sample:", json.dumps({
            "excel_row": item["excel_row"],
            "business_key": _business_key_text(values),
            "id": values.get("recipe_id"),
            "name": values.get("name"),
            "lang": values.get("lang"),
            "reasons": item["reasons"],
            "ingredients": values.get("ingredients"),
            "image_url": values.get("image_url"),
        }, ensure_ascii=False, default=str))

    if args.report:
        _write_report(Path(args.report), invalid_rows)
        print(f"Report: {args.report}")

    output_path = None
    if args.in_place:
        backup = src.with_suffix(src.suffix + ".bak")
        shutil.copy2(src, backup)
        output_path = src
        print(f"Backup: {backup}")
    elif args.out:
        output_path = Path(args.out)

    if output_path:
        for item in sorted(invalid_rows, key=lambda x: x["excel_row"], reverse=True):
            ws.delete_rows(item["excel_row"], 1)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        print(f"Cleaned Excel: {output_path}")
        print(f"Deleted rows from output: {len(invalid_rows)}")
        return 0

    if invalid_rows:
        print("❌ 存在不合规数据；请修正源 Excel，或使用 --out 生成清洗副本后入库。")
        return 1

    print("✅ Excel 校验通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
