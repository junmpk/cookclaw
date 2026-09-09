"""
菜谱标签富化 — 用大模型(Qwen)给每道菜生成「中英双语」多维度标签，
读取源 Excel 的 zh + en 两页，或读取已合并且带 lang 列的单 sheet，
输出一份带 lang + 标签 列的新 Excel。

标签维度：菜系 / 口味 / 烹饪方式 / 主食材类别 / 营养特点 / 适用场景人群 / 餐次。
双语：每个维度尽量给出中文与英文成对标签，最大化跨语言检索召回。

用法（在本技能目录下用独立 venv 运行）：
  # 抽样预览（默认5条，不写文件）
  .venv/bin/python enrich_tags.py --src tmp/recipes_merged_for_ingest.xlsx --limit 5 --dry
  # 全量生成 → 写出富化后的 Excel（原文件不动）
  .venv/bin/python enrich_tags.py --src tmp/recipes_merged_for_ingest.xlsx --out tmp/recipes_merged_enriched_for_ingest.xlsx

特性：异步并发 + 断点缓存 tag_cache.jsonl（键=lang:id），中断可续跑。
需要环境变量：DASHSCOPE_API_KEY（已自动从项目根 .env 读取）。
"""
import os
import sys
import json
import asyncio
import argparse
from pathlib import Path

import httpx
from dotenv import load_dotenv

_PROJECT_ROOT = Path(__file__).resolve().parents[4]
load_dotenv(_PROJECT_ROOT / ".env")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import EMBEDDING_CONFIG  # noqa: E402  复用 api_key / base_url
import openpyxl                       # noqa: E402
from openai import AsyncOpenAI        # noqa: E402
from loguru import logger             # noqa: E402

# ── 配置 ──────────────────────────────────────────────────────
SOURCE_XLSX = _PROJECT_ROOT / "菜品清单_2026.04.08.xlsx"
DEFAULT_OUT = _PROJECT_ROOT / "菜品清单_2026.04.08_enriched.xlsx"
SHEETS = ("zh", "en")
TAG_COL = "标签"
CACHE = Path(__file__).resolve().parent / "tag_cache.jsonl"
MODEL = os.getenv("TAG_MODEL", "qwen-plus")
CONCURRENCY = 8

SYSTEM_PROMPT = """你是中/西餐菜谱标签专家。根据菜名和食材，为这道菜生成「中英双语」标签。
请覆盖以下维度，每个维度尽量同时给出中文标签和对应英文标签（成对出现）：
- 菜系：川菜/Sichuan、粤菜/Cantonese、家常菜/home-style、西餐/Western、日料/Japanese 等
- 口味：咸鲜/savory、香辣/spicy、酸甜/sweet-and-sour、清淡/light 等
- 烹饪方式：炒/stir-fried、炖/stewed、蒸/steamed、烤/roasted、煮/boiled 等
- 主食材类别：猪肉/pork、牛肉/beef、鸡肉/chicken、水产/seafood、蔬菜/vegetable、豆制品/tofu 等
- 营养特点：高蛋白/high-protein、低脂/low-fat、高纤维/high-fiber、补钙/calcium-rich 等
- 适用场景/人群：快手菜/quick、下饭菜/rice-companion、减脂餐/diet、宴客/banquet、儿童/kids 等
- 餐次/类型：早餐/breakfast、正餐/main-course、汤品/soup、甜点/dessert、饮品/beverage 等

只输出一个 JSON 数组，元素是标签字符串（中英混合，成对），共 16-28 个；不要解释、不要代码块。
示例：["川菜","Sichuan cuisine","香辣","spicy","炒","stir-fried","猪肉","pork","高蛋白","high-protein","下饭菜","rice-companion","正餐","main course"]"""

_client = AsyncOpenAI(
    api_key=EMBEDDING_CONFIG["api_key"],
    base_url=EMBEDDING_CONFIG["base_url"],
    http_client=httpx.AsyncClient(trust_env=False),  # 忽略代理，直连 DashScope
)


async def gen_tags(sem, name, ingredients):
    """调用 Qwen 生成双语标签 list；失败重试 3 次，最终失败返回 []。"""
    async with sem:
        for attempt in range(3):
            try:
                resp = await _client.chat.completions.create(
                    model=MODEL,
                    messages=[
                        {"role": "system", "content": SYSTEM_PROMPT},
                        {"role": "user", "content": f"菜名：{name}\n食材：{ingredients}"},
                    ],
                    temperature=0.3,
                    max_tokens=400,
                )
                raw = (resp.choices[0].message.content or "").strip()
                if raw.startswith("```"):
                    raw = raw.split("\n", 1)[-1].rsplit("```", 1)[0].strip()
                tags = json.loads(raw)
                if isinstance(tags, list):
                    return [str(t).strip() for t in tags if str(t).strip()]
            except Exception as e:
                if attempt == 2:
                    logger.warning(f"标签生成失败：error_type={type(e).__name__}")
                await asyncio.sleep(1.5 * (attempt + 1))
        return []


def _cell(row, idx, name, default=""):
    pos = idx.get(name)
    if pos is None:
        return default
    val = row[pos]
    return default if val is None else val


def _row_to_recipe(row, idx, lang):
    return {
        "id": str(_cell(row, idx, "id")).strip(),
        "name": str(_cell(row, idx, "名称")).strip(),
        "image": str(_cell(row, idx, "图片地址")).strip(),
        "ing": str(_cell(row, idx, "食材")).strip(),
        "lang": str(lang).strip().lower(),
    }


def _read_sheet(ws, lang=None):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    h = [str(x).strip() if x is not None else "" for x in rows[0]]
    idx = {k: i for i, k in enumerate(h)}
    missing = {"id", "名称", "食材", "图片地址"} - set(idx)
    if missing:
        raise ValueError(f"sheet {ws.title!r} 缺少列：{', '.join(sorted(missing))}")

    if lang is None and "lang" not in idx:
        raise ValueError(f"sheet {ws.title!r} 缺少 lang 列，且未从 sheet 名推断语言")

    out = []
    for r in rows[1:]:
        recipe_lang = lang if lang is not None else _cell(r, idx, "lang")
        d = _row_to_recipe(r, idx, recipe_lang)
        if d["id"]:
            out.append(d)
    return out


def _read_all(src):
    """读取旧 zh/en 多 sheet 或新单 sheet → 合并的菜谱列表（每条带 lang）。"""
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    out = []
    try:
        if all(sn in wb.sheetnames for sn in SHEETS):
            for sn in SHEETS:
                out.extend(_read_sheet(wb[sn], lang=sn))
        else:
            out.extend(_read_sheet(wb[wb.sheetnames[0]], lang=None))
        return out
    finally:
        wb.close()


def _uid(d):
    return f'{d["lang"]}:{d["id"]}'


def _load_cache():
    cache = {}
    if CACHE.exists():
        for line in CACHE.read_text(encoding="utf-8").splitlines():
            try:
                o = json.loads(line)
                cache[o["uid"]] = o["tags"]
            except Exception:
                pass
    return cache


async def main_async(args):
    if not EMBEDDING_CONFIG.get("api_key") or "填入" in EMBEDDING_CONFIG["api_key"]:
        print("❌ 未读到真实 DASHSCOPE_API_KEY，请先在项目根 .env 填好。")
        sys.exit(1)

    all_rows = _read_all(args.src)
    rows = all_rows[:args.limit] if args.limit else all_rows
    n_zh = sum(1 for d in all_rows if d["lang"] == "zh")
    n_en = sum(1 for d in all_rows if d["lang"] == "en")
    logger.info(f"读取并集 {len(all_rows)} 条（zh={n_zh}, en={n_en}）；本次处理 {len(rows)} 条")

    cache = {} if args.dry else _load_cache()

    sem = asyncio.Semaphore(CONCURRENCY)
    todo = [d for d in rows if _uid(d) not in cache]
    logger.info(f"待生成 {len(todo)} 条（已缓存命中 {len(rows) - len(todo)}）")

    results = {}

    async def worker(d):
        tags = await gen_tags(sem, d["name"], d["ing"])
        results[_uid(d)] = tags
        if not args.dry:
            with open(CACHE, "a", encoding="utf-8") as f:
                f.write(json.dumps({"uid": _uid(d), "name": d["name"], "tags": tags}, ensure_ascii=False) + "\n")
        return d["name"], d["lang"], tags

    done = 0
    tasks = [asyncio.create_task(worker(d)) for d in todo]
    for fut in asyncio.as_completed(tasks):
        name, lang, tags = await fut
        done += 1
        if args.dry or done <= 5 or done % 100 == 0:
            print(f"  [{done}/{len(todo)}] [{lang}] {name} → {tags}")

    if args.dry:
        print("\n（dry 模式：未写文件。确认双语标签 OK 后去掉 --dry/--limit 跑全量。）")
        return

    # ── 写出富化后的新 Excel（单 sheet，并集，保留原文件不动）──
    all_tags = {**cache, **results}
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "recipes"
    out_ws.append(["id", "名称", "图片地址", "食材", "lang", TAG_COL])
    n_tagged = 0
    for d in all_rows:
        tags = all_tags.get(_uid(d), [])
        if tags:
            n_tagged += 1
        out_ws.append([d["id"], d["name"], d["image"], d["ing"], d["lang"], "、".join(tags)])
    out_path = Path(args.out)
    out_wb.save(out_path)
    print(f"\n✅ 已写出富化 Excel：{out_path}")
    print(f"   共 {len(all_rows)} 行（zh={n_zh}, en={n_en}），其中 {n_tagged} 行有标签。")
    print(f"   下一步灌库：.venv/bin/python ingest_local.py \"{out_path}\" --recreate")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(SOURCE_XLSX), help="输入 Excel 路径")
    ap.add_argument("--limit", type=int, default=0, help="只处理前 N 条（抽样用）")
    ap.add_argument("--dry", action="store_true", help="只打印不写文件/缓存")
    ap.add_argument("--out", default=str(DEFAULT_OUT), help="输出 Excel 路径")
    args = ap.parse_args()
    if args.dry and not args.limit:
        args.limit = 5
    asyncio.run(main_async(args))


if __name__ == "__main__":
    main()
