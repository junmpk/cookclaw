#!/usr/bin/env python3
"""用项目 qwen3.7-plus 为标准化菜谱生成单语标签、canonical 与营养估算。

设计原则：
- 模型只选择受控 canonical code，并估算每 100g 营养；展示标签由本地词表映射。
- 营养标签、减脂餐、健身餐、快手菜按确定性规则派生，不允许模型自由发挥。
- 缓存按 ``lang:id`` 逐条追加，任务中断后可直接续跑。
- 原文件不覆盖，输出一份新增「标签标准编码」「营养成分」的新工作簿。
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import time
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import httpx
from dotenv import load_dotenv
from openai import AsyncOpenAI
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")

from app.core.config import settings  # noqa: E402


DISPLAY: dict[str, dict[str, tuple[str, str]]] = {
    "record_type": {
        "device_program": ("设备程序", "Device program"),
    },
    "cuisine": {
        "hunan": ("湘菜", "Hunan"),
        "sichuan": ("川菜", "Sichuan"),
        "cantonese": ("粤菜", "Cantonese"),
        "shandong": ("鲁菜", "Shandong"),
        "jiangsu": ("苏菜", "Jiangsu"),
        "zhejiang": ("浙菜", "Zhejiang"),
        "fujian": ("闽菜", "Fujian"),
        "anhui": ("徽菜", "Anhui"),
        "northeastern_chinese": ("东北菜", "Northeastern Chinese"),
        "northwestern_chinese": ("西北菜", "Northwestern Chinese"),
        "home_style": ("家常菜", "Home-style"),
        "western": ("西餐", "Western"),
        "japanese": ("日料", "Japanese"),
        "korean": ("韩餐", "Korean"),
        "thai": ("泰式", "Thai"),
        "vietnamese": ("越南菜", "Vietnamese"),
        "italian": ("意大利菜", "Italian"),
        "french": ("法餐", "French"),
        "mediterranean": ("地中海菜", "Mediterranean"),
        "mexican": ("墨西哥菜", "Mexican"),
        "indian": ("印度菜", "Indian"),
        "middle_eastern": ("中东菜", "Middle Eastern"),
        "moroccan": ("摩洛哥菜", "Moroccan"),
        "american": ("美式", "American"),
    },
    "flavor": {
        "spicy": ("香辣", "Spicy"),
        "numbing_spicy": ("麻辣", "Numbing-spicy"),
        "savory": ("咸鲜", "Savory"),
        "light": ("清淡", "Light"),
        "sweet": ("甜味", "Sweet"),
        "sweet_sour": ("酸甜", "Sweet-and-sour"),
        "sour_spicy": ("酸辣", "Sour-and-spicy"),
        "sour": ("酸味", "Sour"),
        "salty": ("咸香", "Salty-savory"),
        "umami": ("鲜味", "Umami"),
        "aromatic": ("鲜香", "Aromatic"),
        "creamy": ("奶香", "Creamy"),
    },
    "method": {
        "stir_fried": ("炒", "Stir-fried"),
        "deep_fried": ("油炸", "Deep-fried"),
        "fried": ("煎", "Fried"),
        "steamed": ("蒸", "Steamed"),
        "stewed": ("炖", "Stewed"),
        "braised": ("红烧", "Braised"),
        "boiled": ("煮", "Boiled"),
        "simmered": ("煨", "Simmered"),
        "roasted": ("烤制", "Roasted"),
        "baked": ("烘焙", "Baked"),
        "grilled": ("烧烤", "Grilled"),
        "cold_tossed": ("凉拌", "Cold-tossed"),
        "blended": ("搅打", "Blended"),
        "juiced": ("榨汁", "Juiced"),
        "pickled": ("腌制", "Pickled"),
        "whipping": ("打发", "Whipping"),
        "fermenting": ("发酵", "Fermenting"),
        "pureeing": ("制泥", "Pureeing"),
        "stirring": ("搅拌", "Stirring"),
        "sous_vide": ("低温慢煮", "Sous vide"),
        "reheating": ("加热", "Reheating"),
        "warming": ("保温", "Keep-warm"),
        "peeling": ("削皮", "Peeling"),
        "slicing": ("切片切丝", "Slicing and shredding"),
        "grinding": ("研磨", "Grinding"),
        "kneading": ("和面", "Kneading"),
        "cleaning": ("清洁", "Cleaning"),
    },
    "main_ingredient": {
        "beef": ("牛肉", "Beef"),
        "chicken": ("鸡肉", "Chicken"),
        "pork": ("猪肉", "Pork"),
        "lamb": ("羊肉", "Lamb"),
        "duck": ("鸭肉", "Duck"),
        "fish": ("鱼类", "Fish"),
        "shrimp": ("虾类", "Shrimp"),
        "seafood": ("海鲜", "Seafood"),
        "egg": ("蛋类", "Egg"),
        "tofu": ("豆制品", "Tofu"),
        "vegetable": ("蔬菜", "Vegetable"),
        "mushroom": ("菌菇", "Mushroom"),
        "potato": ("薯类", "Potato"),
        "rice": ("米饭", "Rice"),
        "noodle": ("面食", "Noodle"),
        "dairy": ("乳制品", "Dairy"),
        "fruit": ("水果", "Fruit"),
        "grain": ("谷物", "Grain"),
        "legume": ("豆类", "Legume"),
        "nuts_seeds": ("坚果种子", "Nuts and seeds"),
        "herb_spice": ("香辛料", "Herbs and spices"),
        "water": ("水", "Water"),
        "prepared_food": ("熟食", "Prepared food"),
        "prepared_ingredient": ("预处理食材", "Prepared ingredient"),
    },
    "meal": {
        "breakfast": ("早餐", "Breakfast"),
        "main_course": ("正餐", "Main course"),
        "soup": ("汤品", "Soup"),
        "dessert": ("甜品", "Dessert"),
        "beverage": ("饮品", "Beverage"),
        "snack": ("小吃", "Snack"),
        "staple": ("主食", "Staple"),
        "side_dish": ("配菜", "Side dish"),
        "sauce": ("酱料", "Sauce"),
        "kitchen_program": ("厨房功能", "Kitchen function"),
    },
    "nutrition": {
        "high_protein": ("高蛋白", "High-protein"),
        "low_fat": ("低脂", "Low-fat"),
        "low_calorie": ("低卡", "Low-calorie"),
        "high_fiber": ("高纤维", "High-fiber"),
        "low_carb": ("低碳水", "Low-carb"),
    },
    "scene": {
        "quick": ("快手菜", "Quick"),
        "rice_companion": ("下饭菜", "Rice-companion"),
        "fat_loss": ("减脂餐", "Fat-loss"),
        "fitness": ("健身餐", "Fitness"),
        "child_friendly": ("儿童友好", "Child-friendly"),
        "elderly_friendly": ("老人友好", "Elderly-friendly"),
        "banquet": ("宴客菜", "Banquet"),
        "bento": ("便当", "Bento"),
        "food_preparation": ("备餐功能", "Food preparation"),
        "device_maintenance": ("设备维护", "Device maintenance"),
    },
    "diet": {
        "vegetarian": ("素食", "Vegetarian"),
        "vegan": ("纯素", "Vegan"),
        "gluten_free": ("无麸质", "Gluten-free"),
    },
    "difficulty": {
        "easy": ("简单", "Easy"),
        "moderate": ("中等", "Moderate"),
        "advanced": ("进阶", "Advanced"),
    },
}

DIM_LIMITS = {
    "cuisine": (0, 1),
    "flavor": (1, 2),
    "method": (1, 2),
    "main_ingredient": (1, 2),
    "meal": (1, 1),
    "scene": (0, 4),
    "diet": (0, 2),
}
TAG_ORDER = (
    "cuisine",
    "flavor",
    "method",
    "main_ingredient",
    "meal",
    "nutrition",
    "scene",
    "diet",
)
NUTRIENT_KEYS = (
    "calorie_kcal",
    "protein_g",
    "fat_g",
    "carbohydrate_g",
    "fiber_g",
)
DIFFICULTY = {
    "0": ("简单", "Easy", "easy"),
    "1": ("中等", "Moderate", "moderate"),
    "2": ("进阶", "Advanced", "advanced"),
}

# 这些行是设备预设程序，不是具有确定配方的菜谱。保留它们供设备功能检索，
# 但禁止伪造口味和每 100g 营养。
DEVICE_PROGRAMS: dict[str, tuple[str, list[str], list[str]]] = {
    "2029024142275284994": ("whipping", ["dairy", "egg"], ["food_preparation"]),
    "2029025688148815873": ("boiled", ["egg"], ["food_preparation"]),
    "2029028070895169538": ("fermenting", ["prepared_ingredient"], ["food_preparation"]),
    "2029029600243265537": ("pureeing", ["fruit", "vegetable"], ["food_preparation"]),
    "2029034475823869954": ("steamed", ["prepared_ingredient"], ["food_preparation"]),
    "2029040618382200834": ("stirring", ["prepared_ingredient"], ["food_preparation"]),
    "2029042618943574018": ("fermenting", ["dairy"], ["food_preparation"]),
    "2029044699465158657": ("sous_vide", ["prepared_ingredient"], ["food_preparation"]),
    "2029046117722796034": ("reheating", ["prepared_food"], ["food_preparation"]),
    "2029047161131413506": ("cleaning", ["prepared_ingredient"], ["food_preparation"]),
    "2029048131127250945": ("warming", ["prepared_food"], ["food_preparation"]),
    "2029049320556367873": ("blended", ["fruit", "vegetable"], ["food_preparation"]),
    "2029050406528782337": ("stewed", ["prepared_ingredient"], ["food_preparation"]),
    "2029051876472827906": ("cleaning", ["water"], ["device_maintenance"]),
    "2029060329878753282": ("peeling", ["potato", "vegetable"], ["food_preparation"]),
    "2029061542842109954": ("slicing", ["prepared_ingredient"], ["food_preparation"]),
    "2029062687241056258": ("grinding", ["herb_spice"], ["food_preparation"]),
    "2029064459615637505": ("cleaning", ["water"], ["device_maintenance"]),
    "2029095979239649281": ("kneading", ["grain"], ["food_preparation"]),
}

SYSTEM_PROMPT = """你是严谨的菜谱元数据标注员。请根据真实菜名、食材重量、调料、
描述、预计用时和步骤，为每道菜选择受控 canonical code，并估算每100克成品营养。

只输出 JSON 对象：{"items":[...]}。每个 item 必须包含：
{
  "uid":"输入uid",
  "cuisine":null或一个受控code,
  "flavor":["1-2个受控code"],
  "method":["1-2个受控code"],
  "main_ingredient":["1-2个受控code"],
  "meal":"一个受控code",
  "scene":["0-2个受控code，只可从 rice_companion/child_friendly/elderly_friendly/banquet/bento 选择"],
  "diet":["0-2个受控code"],
  "nutrition":{"calorie_kcal":数值,"protein_g":数值,"fat_g":数值,"carbohydrate_g":数值,"fiber_g":数值}
}

受控 code：
- cuisine: hunan,sichuan,cantonese,shandong,jiangsu,zhejiang,fujian,anhui,
  northeastern_chinese,northwestern_chinese,home_style,western,japanese,korean,
  thai,vietnamese,italian,french,mediterranean,mexican,indian,middle_eastern,
  moroccan,american。菜系证据不足必须为 null，不要用 home_style 兜底。
- flavor: spicy,numbing_spicy,savory,light,sweet,sweet_sour,sour_spicy,sour,
  salty,umami,aromatic,creamy。
- method: stir_fried,deep_fried,fried,steamed,stewed,braised,boiled,simmered,
  roasted,baked,grilled,cold_tossed,blended,juiced,pickled。
  Stir-fried、Deep-fried、Fried 必须严格区分，不能用子串联想。
- main_ingredient: beef,chicken,pork,lamb,duck,fish,shrimp,seafood,egg,tofu,
  vegetable,mushroom,potato,rice,noodle,dairy,fruit,grain,legume,nuts_seeds,
  herb_spice。芝麻、花生、坚果必须归 nuts_seeds；姜及香辛料归 herb_spice，
  不得因为含糖就归 grain。
- meal: breakfast,main_course,soup,dessert,beverage,snack,staple,side_dish,sauce。
- scene: rice_companion,child_friendly,elderly_friendly,banquet,bento。
- diet: vegetarian,vegan,gluten_free。

严格规则：
1. 不输出中文或英文展示标签，只输出 canonical code。
2. cuisine 不确定就填 null，不强行判断。
3. 儿童友好：必须不辣、无酒精/咖啡因、无明显整颗坚果或难处理骨刺。
4. 老人友好：只限蒸、炖、煮、煨、红烧等偏软做法，且不重辣。
5. gluten_free 只有输入原文明确出现“无麸质”或“gluten-free”时才可标，
   不能仅凭食材看似天然无麸质就推断。
6. 不生成糖尿病适用、孕妇适用、降血压等医疗或特殊人群结论。
7. nutrition 是每100克成品合理估算。优先按食材及调料重量估算；重量缺失时结合
   菜名、食材总量和做法估算。五项都必须给数值，不得给 null。热量范围 0-900，
   蛋白质/脂肪/碳水/纤维范围 0-100，保留一位小数即可。
8. 不要输出 high_protein、low_fat、low_calorie、high_fiber、low_carb、quick、
   fat_loss、fitness；这些由程序按阈值派生。
9. 输入有几道菜就返回几条，uid 必须原样返回，不得遗漏或新增。
"""


def clean(value: Any) -> str:
    return re.sub(
        r"\s+",
        " ",
        str("" if value is None else value),
    ).strip()


def parse_steps(value: Any) -> list[str]:
    if isinstance(value, list):
        return [clean(item) for item in value if clean(item)]
    text = str(value or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, list):
        return [clean(item) for item in parsed if clean(item)]
    return [clean(item) for item in re.split(r"\r?\n+", text) if clean(item)]


def extract_json(content: str) -> dict:
    text = str(content or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if not match:
            return {}
        parsed = json.loads(match.group(0))
    return parsed if isinstance(parsed, dict) else {}


def codes(value: Any, dim: str, *, max_count: int) -> list[str]:
    allowed = DISPLAY[dim]
    if isinstance(value, str):
        value = [value]
    result = []
    for raw in value or []:
        code = clean(raw).lower()
        if code in allowed and code not in result:
            result.append(code)
        if len(result) >= max_count:
            break
    return result


def safe_float(value: Any, low: float, high: float) -> float | None:
    try:
        number = round(float(value), 1)
    except (TypeError, ValueError):
        return None
    if not low <= number <= high:
        return None
    return number


def validate_model_item(item: dict, expected_uid: str) -> dict | None:
    if clean(item.get("uid")) != expected_uid:
        return None
    cuisine_raw = clean(item.get("cuisine")).lower()
    cuisine = [cuisine_raw] if cuisine_raw in DISPLAY["cuisine"] else []
    result = {
        "cuisine": cuisine[:1],
        "flavor": codes(item.get("flavor"), "flavor", max_count=2),
        "method": codes(item.get("method"), "method", max_count=2),
        "main_ingredient": codes(
            item.get("main_ingredient"), "main_ingredient", max_count=2
        ),
        "meal": codes(item.get("meal"), "meal", max_count=1),
        "scene": codes(item.get("scene"), "scene", max_count=2),
        "diet": codes(item.get("diet"), "diet", max_count=2),
    }
    if any(
        len(result[dim]) < DIM_LIMITS[dim][0]
        for dim in ("flavor", "method", "main_ingredient", "meal")
    ):
        return None
    raw_nutrition = item.get("nutrition") or {}
    nutrition = {
        "basis": "per_100g_estimated",
        "calorie_kcal": safe_float(raw_nutrition.get("calorie_kcal"), 0, 900),
        "protein_g": safe_float(raw_nutrition.get("protein_g"), 0, 100),
        "fat_g": safe_float(raw_nutrition.get("fat_g"), 0, 100),
        "carbohydrate_g": safe_float(
            raw_nutrition.get("carbohydrate_g"), 0, 100
        ),
        "fiber_g": safe_float(raw_nutrition.get("fiber_g"), 0, 100),
    }
    if any(nutrition[key] is None for key in NUTRIENT_KEYS):
        return None
    result["nutrition_estimate"] = nutrition
    return result


def contains_any(text: str, terms: tuple[str, ...]) -> bool:
    lower = text.lower()
    return any(term.lower() in lower for term in terms)


def derive_metadata(record: dict, model_data: dict) -> dict:
    result = {
        dim: list(model_data.get(dim) or [])
        for dim in ("cuisine", "flavor", "method", "main_ingredient", "meal", "scene", "diet")
    }
    nutrition = dict(model_data["nutrition_estimate"])
    calorie = float(nutrition["calorie_kcal"])
    protein = float(nutrition["protein_g"])
    fat = float(nutrition["fat_g"])
    carbs = float(nutrition["carbohydrate_g"])
    fiber = float(nutrition["fiber_g"])
    method = set(result["method"])
    flavor = set(result["flavor"])
    source = " ".join(
        [
            record["name"],
            record["ingredients"],
            record["seasonings"],
            record["steps_text"],
        ]
    )
    fried = bool(method & {"deep_fried", "fried"})

    animal_flesh_terms = (
        "猪肉", "五花肉", "排骨", "牛肉", "羊肉", "鸡肉", "鸡胸", "鸡腿",
        "鸡翅", "鸭肉", "鱼肉", "鲜鱼", "虾", "螃蟹", "蟹肉", "贝类",
        "蛤蜊", "海参", "鲍鱼", "海鲜",
        "pork", "beef", "lamb", "mutton", "chicken", "duck", "fish",
        "shrimp", "prawn", "crab", "clam", "seafood", "bacon", "ham",
    )
    animal_product_terms = animal_flesh_terms + (
        "蛋", "牛奶", "奶油", "黄油", "芝士", "乳酪", "酸奶", "蜂蜜",
        "egg", "milk", "cream", "butter", "cheese", "yogurt", "honey",
    )
    diet = list(result["diet"])
    if contains_any(source, animal_flesh_terms):
        diet = [code for code in diet if code not in {"vegetarian", "vegan"}]
    elif contains_any(source, animal_product_terms):
        diet = [code for code in diet if code != "vegan"]
    if not contains_any(source, ("无麸质", "gluten-free", "gluten free")):
        diet = [code for code in diet if code != "gluten_free"]
    if "vegan" in diet:
        diet = [code for code in diet if code != "vegetarian"]
    result["diet"] = list(dict.fromkeys(diet))[:2]

    nutrition_codes = []
    if protein >= 10 or (calorie > 0 and protein * 4 / calorie >= 0.20):
        nutrition_codes.append("high_protein")
    if fat <= 3 and not fried:
        nutrition_codes.append("low_fat")
    if calorie <= 120:
        nutrition_codes.append("low_calorie")
    if fiber >= 3:
        nutrition_codes.append("high_fiber")
    if carbs <= 10:
        nutrition_codes.append("low_carb")
    result["nutrition"] = nutrition_codes

    scenes = [
        code
        for code in result["scene"]
        if code not in {"quick", "fat_loss", "fitness"}
    ]
    time_min = record.get("time_min")
    step_count = int(record.get("step_count") or 0)
    if isinstance(time_min, (int, float)) and time_min <= 20 and step_count <= 6:
        scenes.append("quick")
    if "low_calorie" in nutrition_codes and any(
        code in nutrition_codes for code in ("low_fat", "high_protein", "high_fiber")
    ):
        scenes.append("fat_loss")
    high_sugar_style = "sweet" in flavor and carbs > 20
    if "high_protein" in nutrition_codes and not fried and not high_sugar_style:
        scenes.append("fitness")

    unsafe_child = (
        bool(flavor & {"spicy", "numbing_spicy", "sour_spicy"})
        or contains_any(
            source,
            (
                "料酒", "黄酒", "白酒", "葡萄酒", "咖啡", "咖啡因", "整颗坚果",
                "骨刺", "cooking wine", "wine", "brandy", "rum", "coffee",
                "caffeine", "whole nuts", "fish bones",
            ),
        )
    )
    if unsafe_child:
        scenes = [code for code in scenes if code != "child_friendly"]
    soft_methods = {"steamed", "stewed", "boiled", "simmered", "braised"}
    if (
        "elderly_friendly" in scenes
        and (
            not method.intersection(soft_methods)
            or bool(flavor & {"spicy", "numbing_spicy", "sour_spicy"})
        )
    ):
        scenes.remove("elderly_friendly")
    result["scene"] = list(dict.fromkeys(scenes))[:4]

    difficulty_raw = clean(record.get("difficulty"))
    difficulty = DIFFICULTY.get(difficulty_raw, DIFFICULTY["1"])
    result["difficulty"] = [difficulty[2]]
    result["nutrition_estimate"] = nutrition
    return result


def device_program_metadata(record: dict) -> dict | None:
    specification = DEVICE_PROGRAMS.get(record["id"])
    if not specification:
        return None
    method, ingredients, scenes = specification
    time_min = record.get("time_min")
    if (
        isinstance(time_min, (int, float))
        and time_min <= 20
        and int(record.get("step_count") or 0) <= 6
    ):
        scenes = [*scenes, "quick"]
    difficulty_raw = clean(record.get("difficulty"))
    difficulty = DIFFICULTY.get(difficulty_raw, DIFFICULTY["1"])
    return {
        "record_type": ["device_program"],
        "cuisine": [],
        "flavor": [],
        "method": [method],
        "main_ingredient": ingredients[:2],
        "meal": ["kitchen_program"],
        "nutrition": [],
        "scene": list(dict.fromkeys(scenes))[:2],
        "diet": [],
        "difficulty": [difficulty[2]],
        "nutrition_estimate": {"basis": "not_applicable"},
    }


def display_tags(metadata: dict, lang: str) -> list[str]:
    language_index = 1 if lang == "en" else 0
    tags = []
    dimensions = (
        ("record_type", "method", "main_ingredient", "meal", "scene", "difficulty")
        if "device_program" in (metadata.get("record_type") or [])
        else TAG_ORDER
    )
    for dim in dimensions:
        for code in metadata.get(dim) or []:
            pair = DISPLAY.get(dim, {}).get(code)
            if pair:
                label = pair[language_index]
                if label not in tags:
                    tags.append(label)
    if "device_program" not in (metadata.get("record_type") or []) and len(tags) < 5:
        for code in metadata.get("difficulty") or []:
            pair = DISPLAY["difficulty"].get(code)
            if pair:
                label = pair[language_index]
                if label not in tags:
                    tags.append(label)
    # 保留核心维度优先级；场景/饮食属性在超过 12 个时才被截断。
    return tags[:12]


def cache_load(path: Path) -> dict[str, dict]:
    result: dict[str, dict] = {}
    if not path.exists():
        return result
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            payload = json.loads(line)
            uid = clean(payload.get("uid"))
            data = payload.get("data")
            if uid and isinstance(data, dict):
                result[uid] = data
        except (json.JSONDecodeError, TypeError):
            continue
    return result


@dataclass
class GenerationConfig:
    model: str
    batch_size: int
    concurrency: int
    max_tokens: int


def record_payload(record: dict) -> dict:
    steps = record["steps"]
    return {
        "uid": record["uid"],
        "lang": record["lang"],
        "name": record["name"],
        "ingredients": record["ingredients"][:3500],
        "seasonings": record["seasonings"][:1800],
        "description": record["description"][:1200],
        "estimated_time_minutes": record.get("time_min"),
        "step_count": record["step_count"],
        "steps": steps[:14],
    }


async def generate_batch(
    client: AsyncOpenAI,
    semaphore: asyncio.Semaphore,
    config: GenerationConfig,
    batch: list[dict],
) -> dict[str, dict]:
    expected = {record["uid"] for record in batch}
    payload = {"recipes": [record_payload(record) for record in batch]}
    async with semaphore:
        for attempt in range(3):
            try:
                response = await client.chat.completions.create(
                    model=config.model,
                    messages=[
                        {"role": "system", "content": SYSTEM_PROMPT},
                        {
                            "role": "user",
                            "content": json.dumps(payload, ensure_ascii=False),
                        },
                    ],
                    temperature=0.1,
                    max_tokens=config.max_tokens,
                    response_format={"type": "json_object"},
                    extra_body={"enable_thinking": False},
                )
                parsed = extract_json(response.choices[0].message.content or "")
                generated: dict[str, dict] = {}
                for item in parsed.get("items") or []:
                    if not isinstance(item, dict):
                        continue
                    uid = clean(item.get("uid"))
                    if uid not in expected:
                        continue
                    valid = validate_model_item(item, uid)
                    if valid:
                        generated[uid] = valid
                if generated:
                    return generated
            except Exception as exc:
                if attempt == 2:
                    print(
                        f"GENERATION batch_failed size={len(batch)} "
                        f"error_type={type(exc).__name__}",
                        flush=True,
                    )
                await asyncio.sleep(1.5 * (attempt + 1))
    return {}


def read_records(
    source: Path,
    limit: int = 0,
    language: str = "",
) -> tuple[list[dict], list[str]]:
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [clean(cell.value) for cell in ws[1]]
    index = {header: position for position, header in enumerate(headers)}
    required = {
        "id", "名称", "食材", "标签", "lang", "菜谱描述", "预计用时",
        "难易程度", "调料", "烹饪步骤",
    }
    missing = sorted(required - set(index))
    if missing:
        wb.close()
        raise RuntimeError(f"源文件缺少字段：{missing}")
    records = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        recipe_id = clean(values[index["id"]])
        if not recipe_id:
            continue
        lang = "en" if clean(values[index["lang"]]).lower() == "en" else "zh"
        if language and lang != language:
            continue
        steps = parse_steps(values[index["烹饪步骤"]])
        time_raw = values[index["预计用时"]]
        try:
            time_min = float(time_raw) if time_raw not in (None, "") else None
        except (TypeError, ValueError):
            time_min = None
        records.append(
            {
                "uid": f"{lang}:{recipe_id}",
                "id": recipe_id,
                "lang": lang,
                "name": clean(values[index["名称"]]),
                "ingredients": clean(values[index["食材"]]),
                "seasonings": clean(values[index["调料"]]),
                "description": clean(values[index["菜谱描述"]]),
                "time_min": time_min,
                "difficulty": clean(values[index["难易程度"]]),
                "steps": steps,
                "steps_text": " ".join(steps),
                "step_count": len(steps),
            }
        )
        if limit and len(records) >= limit:
            break
    wb.close()
    return records, headers


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def write_workbook(
    source: Path,
    output: Path,
    generated: dict[str, dict],
) -> dict:
    wb = load_workbook(source)
    ws = wb[wb.sheetnames[0]]
    headers = [clean(cell.value) for cell in ws[1]]
    original_index = {header: position + 1 for position, header in enumerate(headers)}
    tags_column = original_index["标签"]
    ws.insert_cols(tags_column + 1, amount=2)
    canonical_column = tags_column + 1
    nutrition_column = tags_column + 2
    ws.cell(1, canonical_column).value = "标签标准编码"
    ws.cell(1, nutrition_column).value = "营养成分"
    copy_cell_style(ws.cell(1, tags_column), ws.cell(1, canonical_column))
    copy_cell_style(ws.cell(1, tags_column), ws.cell(1, nutrition_column))

    headers = [clean(cell.value) for cell in ws[1]]
    index = {header: position + 1 for position, header in enumerate(headers)}
    language_counts = {"zh": 0, "en": 0}
    tag_counts = []
    for row in range(2, ws.max_row + 1):
        recipe_id = clean(ws.cell(row, index["id"]).value)
        lang = "en" if clean(ws.cell(row, index["lang"]).value).lower() == "en" else "zh"
        uid = f"{lang}:{recipe_id}"
        data = generated.get(uid)
        if not data:
            raise RuntimeError(f"缺少生成缓存：{uid}")
        tags = display_tags(data, lang)
        if not 5 <= len(tags) <= 12:
            raise RuntimeError(f"标签数量异常：{uid} count={len(tags)}")
        language_counts[lang] += 1
        tag_counts.append(len(tags))
        ws.cell(row, index["标签"]).value = "、".join(tags)
        ws.cell(row, canonical_column).value = json.dumps(
            {
                dim: data.get(dim) or []
                for dim in (
                    "record_type", "cuisine", "flavor", "method", "main_ingredient", "meal",
                    "nutrition", "scene", "diet", "difficulty",
                )
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )
        ws.cell(row, nutrition_column).value = json.dumps(
            data["nutrition_estimate"],
            ensure_ascii=False,
            separators=(",", ":"),
        )
        difficulty_code = (data.get("difficulty") or ["moderate"])[0]
        difficulty_pair = next(
            (pair for pair in DIFFICULTY.values() if pair[2] == difficulty_code),
            DIFFICULTY["1"],
        )
        ws.cell(row, index["难易程度"]).value = (
            difficulty_pair[1] if lang == "en" else difficulty_pair[0]
        )
        copy_cell_style(ws.cell(row, tags_column), ws.cell(row, canonical_column))
        copy_cell_style(ws.cell(row, tags_column), ws.cell(row, nutrition_column))

    ws.column_dimensions[get_column_letter(tags_column)].width = 38
    ws.column_dimensions[get_column_letter(canonical_column)].width = 52
    ws.column_dimensions[get_column_letter(nutrition_column)].width = 42
    ws.auto_filter.ref = ws.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    return {
        "rows": len(tag_counts),
        "languages": language_counts,
        "tag_min": min(tag_counts),
        "tag_max": max(tag_counts),
        "tag_average": round(sum(tag_counts) / len(tag_counts), 2),
    }


def validate_output(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [clean(cell.value) for cell in ws[1]]
    index = {header: position for position, header in enumerate(headers)}
    required = {"标签", "标签标准编码", "营养成分", "难易程度", "lang"}
    if not required.issubset(index):
        wb.close()
        raise RuntimeError(f"输出字段不完整：{headers}")
    errors = []
    language_counts = {"zh": 0, "en": 0}
    tag_counts = []
    for excel_row, values in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        lang = "en" if clean(values[index["lang"]]).lower() == "en" else "zh"
        language_counts[lang] += 1
        tags = [item for item in clean(values[index["标签"]]).split("、") if item]
        tag_counts.append(len(tags))
        if not 5 <= len(tags) <= 12:
            errors.append((excel_row, "tag_count", len(tags)))
        if lang == "zh" and any(re.search(r"[A-Za-z]", tag) for tag in tags):
            errors.append((excel_row, "zh_has_latin", tags))
        if lang == "en" and any(re.search(r"[\u4e00-\u9fff]", tag) for tag in tags):
            errors.append((excel_row, "en_has_cjk", tags))
        try:
            canonical = json.loads(clean(values[index["标签标准编码"]]))
            nutrition = json.loads(clean(values[index["营养成分"]]))
        except json.JSONDecodeError:
            errors.append((excel_row, "invalid_json", None))
            continue
        is_device_program = "device_program" in (canonical.get("record_type") or [])
        expected_basis = (
            "not_applicable" if is_device_program else "per_100g_estimated"
        )
        if nutrition.get("basis") != expected_basis:
            errors.append((excel_row, "nutrition_basis", nutrition))
        if not is_device_program and any(
            key not in nutrition for key in NUTRIENT_KEYS
        ):
            errors.append((excel_row, "nutrition_fields", nutrition))
        if (
            not canonical.get("method")
            or (not is_device_program and not canonical.get("flavor"))
        ):
            errors.append((excel_row, "canonical_required", canonical))
        expected_difficulty = {"Easy", "Moderate", "Advanced"} if lang == "en" else {
            "简单", "中等", "进阶",
        }
        if clean(values[index["难易程度"]]) not in expected_difficulty:
            errors.append(
                (excel_row, "difficulty_language", values[index["难易程度"]])
            )
        if len(errors) >= 100:
            break
    wb.close()
    if errors:
        raise RuntimeError(
            "输出校验失败：" + json.dumps(errors[:20], ensure_ascii=False)
        )
    return {
        "rows": sum(language_counts.values()),
        "languages": language_counts,
        "headers": headers,
        "tag_min": min(tag_counts),
        "tag_max": max(tag_counts),
        "tag_average": round(sum(tag_counts) / len(tag_counts), 2),
        "errors": 0,
    }


async def run(args: argparse.Namespace) -> int:
    started = time.monotonic()
    records, _ = read_records(
        args.input,
        limit=args.limit,
        language=args.language,
    )
    if settings.LLM_MODEL != "qwen3.7-plus" and not args.allow_model_override:
        raise RuntimeError(
            f"项目主模型不是 qwen3.7-plus：{settings.LLM_MODEL!r}"
        )
    model = settings.LLM_MODEL
    api_key = clean(settings.DASHSCOPE_API_KEY)
    if not api_key:
        raise RuntimeError("项目 .env 缺少 DASHSCOPE_API_KEY")
    cache = cache_load(args.cache)
    records_by_uid = {record["uid"]: record for record in records}
    for record in records:
        difficulty = DIFFICULTY.get(
            clean(record.get("difficulty")),
            DIFFICULTY["1"],
        )
        if record["uid"] in cache:
            cache[record["uid"]]["difficulty"] = [difficulty[2]]
        device_data = device_program_metadata(record)
        if device_data:
            cache[record["uid"]] = device_data
    pending = [record for record in records if record["uid"] not in cache]
    print(
        f"START model={model} rows={len(records)} cache_hits={len(records)-len(pending)} "
        f"pending={len(pending)} batch_size={args.batch_size} concurrency={args.concurrency}",
        flush=True,
    )
    config = GenerationConfig(
        model=model,
        batch_size=args.batch_size,
        concurrency=args.concurrency,
        max_tokens=args.max_tokens,
    )
    client = AsyncOpenAI(
        api_key=api_key,
        base_url=settings.DASHSCOPE_BASE_URL,
        http_client=httpx.AsyncClient(trust_env=False, timeout=120),
    )
    semaphore = asyncio.Semaphore(max(1, args.concurrency))
    batches = [
        pending[index:index + args.batch_size]
        for index in range(0, len(pending), args.batch_size)
    ]
    completed_rows = 0
    tasks = [
        asyncio.create_task(generate_batch(client, semaphore, config, batch))
        for batch in batches
    ]
    args.cache.parent.mkdir(parents=True, exist_ok=True)
    for completed_batches, future in enumerate(asyncio.as_completed(tasks), start=1):
        result = await future
        for uid, raw_data in result.items():
            record = records_by_uid[uid]
            data = derive_metadata(record, raw_data)
            cache[uid] = data
            with args.cache.open("a", encoding="utf-8") as handle:
                handle.write(
                    json.dumps(
                        {"uid": uid, "model": model, "data": data},
                        ensure_ascii=False,
                    )
                    + "\n"
                )
        completed_rows += len(result)
        if (
            completed_batches <= 3
            or completed_batches % 20 == 0
            or completed_batches == len(batches)
        ):
            print(
                f"PROGRESS batches={completed_batches}/{len(batches)} "
                f"generated_rows={completed_rows}/{len(pending)} "
                f"cache_total={len(cache)}",
                flush=True,
            )
    await client.close()

    missing = [record for record in records if record["uid"] not in cache]
    if missing:
        print(
            f"RETRY missing={len(missing)} as individual requests",
            flush=True,
        )
        retry_client = AsyncOpenAI(
            api_key=api_key,
            base_url=settings.DASHSCOPE_BASE_URL,
            http_client=httpx.AsyncClient(trust_env=False, timeout=120),
        )
        retry_semaphore = asyncio.Semaphore(max(1, args.concurrency))
        retry_tasks = [
            asyncio.create_task(
                generate_batch(retry_client, retry_semaphore, config, [record])
            )
            for record in missing
        ]
        for future in asyncio.as_completed(retry_tasks):
            result = await future
            for uid, raw_data in result.items():
                record = records_by_uid[uid]
                data = derive_metadata(record, raw_data)
                cache[uid] = data
                with args.cache.open("a", encoding="utf-8") as handle:
                    handle.write(
                        json.dumps(
                            {"uid": uid, "model": model, "data": data},
                            ensure_ascii=False,
                        )
                        + "\n"
                    )
        await retry_client.close()

    missing = [record["uid"] for record in records if record["uid"] not in cache]
    if missing:
        raise RuntimeError(
            f"仍有 {len(missing)} 条生成失败，保留缓存后停止写 Excel：{missing[:20]}"
        )
    if args.dry_run:
        for record in records[: min(10, len(records))]:
            data = cache[record["uid"]]
            print(
                json.dumps(
                    {
                        "uid": record["uid"],
                        "name": record["name"],
                        "tags": display_tags(data, record["lang"]),
                        "canonical": {
                            dim: data.get(dim) or []
                            for dim in TAG_ORDER
                        },
                        "nutrition": data["nutrition_estimate"],
                    },
                    ensure_ascii=False,
                )
            )
        return 0

    write_summary = write_workbook(args.input, args.output, cache)
    validation = validate_output(args.output)
    audit = {
        "source": str(args.input),
        "output": str(args.output),
        "cache": str(args.cache),
        "model": model,
        "elapsed_seconds": round(time.monotonic() - started, 2),
        "write_summary": write_summary,
        "validation": validation,
    }
    args.audit.parent.mkdir(parents=True, exist_ok=True)
    args.audit.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2), flush=True)
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--cache", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--language", choices=("zh", "en"), default="")
    parser.add_argument("--batch-size", type=int, default=8)
    parser.add_argument("--concurrency", type=int, default=6)
    parser.add_argument("--max-tokens", type=int, default=7000)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--allow-model-override", action="store_true")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(asyncio.run(run(parse_args())))
