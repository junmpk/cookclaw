#!/usr/bin/env python3
"""将 Milvus 正式菜谱及全部非空 ``recipe_detail`` 字段导出为 Excel。

保留参考文件的 6 个基础字段；``recipe_detail`` 顶层字段逐列导出，嵌套结构
保存为 JSON。只移除在全部菜谱中为空的详情字段，以及 JSON 内部的空值。
不会导出 dense/sparse 向量、Milvus 内部主键或其它 metadata 运行字段。
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from copy import copy
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = PROJECT_ROOT / "app/agent/skills/recipe-search"
load_dotenv(PROJECT_ROOT / ".env")
sys.path.insert(0, str(SKILL_ROOT))
os.environ.setdefault("MILVUS_ANALYZER_TYPE", "standard")

from dataset.milvus import build_milvus_client, safe_milvus_uri  # noqa: E402


BASE_HEADERS = [
    "id",
    "名称",
    "图片地址",
    "食材",
    "lang",
    "标签",
]
BASE_WIDTHS = {
    "id": 23,
    "名称": 30,
    "图片地址": 58,
    "食材": 54,
    "lang": 10,
    "标签": 52,
}
DETAIL_ORDER = [
    "recipe_id",
    "cookId",
    "language",
    "name",
    "introduction",
    "tips",
    "media",
    "tags",
    "ingredients",
    "steps",
    "cooking_time_seconds",
    "servings",
    "challenge_level",
    "calorie_number",
    "category_ids",
    "accessory_ids",
    "device_model_ids",
    "is_custom_food",
    "executable",
    "source_created_at",
    "source_updated_at",
    "ai_generated",
    "generated_fields",
    "generation_model",
]
QUERY_LIMIT = 16_384
EXCEL_CELL_LIMIT = 32_767


def _args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--template",
        type=Path,
        required=True,
        help="只参考工作表与基础列结构的 Excel",
    )
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="导出目标，可与 template 相同",
    )
    parser.add_argument("--collection", default="recipe_hybrid")
    return parser.parse_args()


def _metadata(value: Any) -> dict:
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return {}
        return parsed if isinstance(parsed, dict) else {}
    return {}


def _text(value: Any) -> str:
    return " ".join(str(value or "").split()).strip()


def _prune_empty(value: Any) -> Any:
    """递归移除 JSON 中的 null、空串、空数组和空对象，保留 0/False。"""
    if isinstance(value, dict):
        result = {}
        for key, raw in value.items():
            cleaned = _prune_empty(raw)
            if cleaned not in (None, "", [], {}):
                result[key] = cleaned
        return result
    if isinstance(value, list):
        return [
            cleaned
            for item in value
            if (cleaned := _prune_empty(item)) not in (None, "", [], {})
        ]
    return value


def _amount_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _ingredient_lines(detail: dict) -> str:
    lines: list[str] = []
    for raw in detail.get("ingredients") or []:
        if not isinstance(raw, dict):
            continue
        name = _text(raw.get("name"))
        amount = _amount_text(raw.get("amount"))
        unit = _text(raw.get("unit"))
        if not name:
            continue
        line = " ".join(part for part in (amount, unit, name) if part)
        remark = _text(raw.get("remark"))
        if remark:
            line += f"（{remark}）"
        lines.append(line)
    return "\n".join(lines)


def _json_cell(value: Any, *, recipe_id: str, field: str) -> str:
    encoded = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if len(encoded) > EXCEL_CELL_LIMIT:
        raise ValueError(
            f"{field}超过 Excel 单元格上限：recipe_id={recipe_id}"
        )
    return encoded


def _tags(value: Any) -> str:
    if not isinstance(value, list):
        return _text(value)
    return "、".join(
        text
        for text in (_text(item) for item in value)
        if text
    )


def _record(metadata: dict) -> dict[str, Any]:
    detail = (
        metadata.get("recipe_detail")
        if isinstance(metadata.get("recipe_detail"), dict)
        else {}
    )
    media = detail.get("media") if isinstance(detail.get("media"), dict) else {}
    recipe_id = _text(
        detail.get("recipe_id")
        or detail.get("cookId")
        or metadata.get("recipe_id")
    )
    name = _text(detail.get("name") or metadata.get("name"))
    image = _text(
        media.get("landscape_image_url")
        or media.get("portrait_image_url")
        or metadata.get("image")
        or metadata.get("image_url")
    )
    lang = "en" if metadata.get("lang") == "en" else "zh"
    ingredients = _ingredient_lines(detail)
    record: dict[str, Any] = {
        "id": recipe_id,
        "名称": name,
        "图片地址": image,
        "食材": ingredients,
        "lang": lang,
        "标签": _tags(detail.get("tags") or metadata.get("tags")),
    }
    for key, raw in detail.items():
        cleaned = _prune_empty(raw)
        if cleaned in (None, "", [], {}):
            continue
        column = f"详情.{key}"
        record[column] = (
            _json_cell(cleaned, recipe_id=recipe_id, field=column)
            if isinstance(cleaned, (dict, list))
            else cleaned
        )
    return record


def _fetch_records(collection: str) -> list[dict[str, Any]]:
    uri = os.getenv("RECIPE_MILVUS_URI", "").strip()
    if not uri:
        raise RuntimeError("RECIPE_MILVUS_URI 未配置")
    client = build_milvus_client(uri)
    try:
        client.load_collection(collection)
        values = client.query(
            collection_name=collection,
            filter="id >= 0",
            output_fields=["metadata"],
            limit=QUERY_LIMIT,
        )
    finally:
        client.close()
    if len(values) >= QUERY_LIMIT:
        raise RuntimeError(f"集合达到查询上限：{QUERY_LIMIT}")
    records = [_record(_metadata(item.get("metadata"))) for item in values]
    records.sort(
        key=lambda record: (
            0 if record["lang"] == "en" else 1,
            str(record["名称"]).casefold(),
            str(record["id"]),
        )
    )
    print(
        f"[fetch] milvus={safe_milvus_uri(uri)} "
        f"collection={collection} rows={len(records)}",
        flush=True,
    )
    return records


def _validate(records: list[dict[str, Any]]) -> dict:
    keys = [(str(row["id"]), str(row["lang"])) for row in records]
    missing_id = sum(not key[0] for key in keys)
    missing_name = sum(not _text(row["名称"]) for row in records)
    missing_ingredients = sum(not _text(row["食材"]) for row in records)
    missing_detail_ingredients = 0
    missing_steps = 0
    invalid_detail = 0
    invalid_structured_quantities = 0
    for row in records:
        try:
            ingredients = json.loads(row["详情.ingredients"])
            steps = json.loads(row["详情.steps"])
        except (TypeError, json.JSONDecodeError):
            invalid_detail += 1
            continue
        if not isinstance(ingredients, list) or not isinstance(steps, list):
            invalid_detail += 1
            continue
        if not ingredients:
            missing_detail_ingredients += 1
        if not steps:
            missing_steps += 1
        invalid_structured_quantities += sum(
            not isinstance(item, dict)
            or item.get("amount") in (None, "", 0, "0")
            or not _text(item.get("unit"))
            for item in ingredients
        )
    invalid_quantity_lines = 0
    for row in records:
        for line in str(row["食材"] or "").splitlines():
            # 导出来源已经通过全量 amount/unit 校验；这里再防止序列化时丢字段。
            if len(line.split(maxsplit=2)) < 3:
                invalid_quantity_lines += 1
    result = {
        "rows": len(records),
        "unique_keys": len(set(keys)),
        "missing_id": missing_id,
        "missing_name": missing_name,
        "missing_ingredients": missing_ingredients,
        "missing_detail_ingredients": missing_detail_ingredients,
        "missing_steps": missing_steps,
        "invalid_detail_json": invalid_detail,
        "invalid_structured_quantities": invalid_structured_quantities,
        "invalid_quantity_lines": invalid_quantity_lines,
    }
    if (
        not records
        or len(keys) != len(set(keys))
        or missing_id
        or missing_name
        or missing_ingredients
        or missing_detail_ingredients
        or missing_steps
        or invalid_detail
        or invalid_structured_quantities
        or invalid_quantity_lines
    ):
        raise RuntimeError(
            f"导出数据校验失败：{json.dumps(result, ensure_ascii=False)}"
        )
    return result


def _columns_and_rows(
    records: list[dict[str, Any]],
) -> tuple[list[str], list[float], list[list[Any]]]:
    present = {
        key.removeprefix("详情.")
        for record in records
        for key in record
        if key.startswith("详情.")
    }
    ordered_detail = [
        field for field in DETAIL_ORDER if field in present
    ]
    ordered_detail.extend(sorted(present - set(ordered_detail)))
    headers = BASE_HEADERS + [f"详情.{field}" for field in ordered_detail]
    widths = []
    for header in headers:
        if header in BASE_WIDTHS:
            widths.append(BASE_WIDTHS[header])
        elif header in {"详情.ingredients", "详情.steps"}:
            widths.append(100)
        elif header in {"详情.media", "详情.tags", "详情.generated_fields"}:
            widths.append(65)
        else:
            widths.append(28)
    rows = [[record.get(header) for header in headers] for record in records]
    return headers, widths, rows


def _write_workbook(
    template: Path,
    output: Path,
    headers: list[str],
    widths: list[float],
    rows: list[list[Any]],
) -> None:
    workbook = load_workbook(template)
    sheet = workbook["recipes"] if "recipes" in workbook.sheetnames else workbook.active

    # 参考文件只提供基础字段；清空旧内容后写入正式集合全量数据。
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    if sheet.max_column > len(headers):
        sheet.delete_cols(len(headers) + 1, sheet.max_column - len(headers))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="1F2937")
    border = Border(bottom=Side(style="thin", color="D9E2F3"))

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = copy(border)
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]

    for row_number, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.font = copy(body_font)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=headers[column - 1] not in {
                    "id",
                    "lang",
                    "详情.cooking_time_seconds",
                    "详情.servings",
                    "详情.challenge_level",
                    "详情.calorie_number",
                },
            )
            cell.border = copy(border)
            if headers[column - 1] in {
                "详情.cooking_time_seconds",
                "详情.servings",
                "详情.challenge_level",
                "详情.calorie_number",
            }:
                cell.number_format = "0.##"
            elif headers[column - 1] == "id":
                cell.number_format = "@"
        sheet.row_dimensions[row_number].height = 48

    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    sheet.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, output)


def main() -> int:
    args = _args()
    if not args.template.exists():
        raise FileNotFoundError(args.template)
    records = _fetch_records(args.collection)
    validation = _validate(records)
    print(f"[validate] {json.dumps(validation, ensure_ascii=False)}", flush=True)
    headers, widths, rows = _columns_and_rows(records)
    print(
        f"[columns] kept={len(headers)} "
        f"detail_fields={json.dumps(headers[len(BASE_HEADERS):], ensure_ascii=False)}",
        flush=True,
    )
    _write_workbook(args.template, args.output, headers, widths, rows)
    print(f"[done] output={args.output} size={args.output.stat().st_size}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
