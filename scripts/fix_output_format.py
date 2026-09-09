"""后处理：修正 AI 生成字段的格式，使其匹配标准文档。

修正项：
1. 调料：JSON 数组 → 换行分隔纯文本
2. 食材：JSON 数组 → 换行分隔纯文本
3. 预计用时：字符串 → 数字

用法：
  python fix_output_format.py input.xlsx -o output.xlsx
"""
import argparse
import json
import openpyxl


def fix_formats(input_path, output_path):
    wb = openpyxl.load_workbook(input_path)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    # Find column indices
    col_map = {h: i for i, h in enumerate(headers)}
    fixes = {"调料": 0, "食材": 0, "预计用时": 0}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for field in ["调料", "食材"]:
            if field in col_map:
                cell = row[col_map[field]]
                val = cell.value
                if val and isinstance(val, str):
                    try:
                        parsed = json.loads(val)
                        if isinstance(parsed, list):
                            cell.value = "\n".join(str(item) for item in parsed)
                            fixes[field] += 1
                    except (json.JSONDecodeError, TypeError):
                        pass

        if "预计用时" in col_map:
            cell = row[col_map["预计用时"]]
            val = cell.value
            if val is not None:
                try:
                    cell.value = int(float(val))
                    fixes["预计用时"] += 1
                except (ValueError, TypeError):
                    pass

    wb.save(output_path)
    print(f"格式修正完成:")
    print(f"  调料: {fixes['调料']} 行从 JSON 数组转为纯文本")
    print(f"  食材: {fixes['食材']} 行从 JSON 数组转为纯文本")
    print(f"  预计用时: {fixes['预计用时']} 行转为数字")
    print(f"  输出: {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args()
    fix_formats(args.input, args.output)
