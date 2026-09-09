#!/usr/bin/env python3
"""导出 Milvus 菜谱集合为 Excel。

默认导出线上检索集合 recipe_hybrid 的可读字段：
Milvus 主键、业务主键、metadata、BM25 index_text。

不导出 dense/sparse 向量列。1024 维向量展开到 Excel 后不可读且文件巨大；
如需向量备份，应另行导出为 parquet/jsonl。
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "app" / "agent" / "skills" / "recipe-search"
load_dotenv(ROOT / ".env")

# 本地 SSH 隧道 / 内网 Milvus 不走系统代理。
_NO_PROXY = "127.0.0.1,localhost,::1"
os.environ["NO_PROXY"] = ",".join(filter(None, [os.environ.get("NO_PROXY", ""), _NO_PROXY]))
os.environ["no_proxy"] = ",".join(filter(None, [os.environ.get("no_proxy", ""), _NO_PROXY]))

sys.path.insert(0, str(SKILL))
from config import MILVUS_CONFIG  # noqa: E402
from dataset.milvus import build_milvus_client  # noqa: E402


HEADERS = [
    "milvus_pk",
    "business_key",
    "recipe_id",
    "lang",
    "name",
    "image_url",
    "ingredients",
    "ingredients_raw",
    "tags",
    "facets",
    "description",
    "index_text",
]


def _metadata(row: dict) -> dict:
    md = row.get("metadata", {})
    if isinstance(md, str):
        try:
            md = json.loads(md)
        except Exception:
            return {}
    return md if isinstance(md, dict) else {}


def _json_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _list_text(value) -> str:
    if isinstance(value, list):
        return "、".join(str(v) for v in value if str(v).strip())
    if value is None:
        return ""
    return str(value)


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in {"", "nan", "none", "null", "未知", "未知食材", "n/a", "na"}
    if isinstance(value, list):
        return not any(not _is_blank(v) for v in value)
    return False


def _valid_url(value) -> bool:
    if _is_blank(value):
        return False
    parsed = urlparse(str(value).strip())
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def _iter_rows(client, collection: str, batch_size: int):
    fields = ["id", "text", "metadata"]
    if hasattr(client, "query_iterator"):
        iterator = client.query_iterator(
            collection_name=collection,
            filter="",
            output_fields=fields,
            batch_size=batch_size,
        )
        try:
            while True:
                batch = iterator.next()
                if not batch:
                    break
                yield from batch
        finally:
            try:
                iterator.close()
            except Exception:
                pass
        return

    offset = 0
    while True:
        batch = client.query(
            collection_name=collection,
            filter="",
            output_fields=fields,
            limit=batch_size,
            offset=offset,
        )
        if not batch:
            break
        yield from batch
        if len(batch) < batch_size:
            break
        offset += len(batch)


def _normalize(row: dict) -> dict:
    md = _metadata(row)
    recipe_id = str(md.get("recipe_id") or "").strip()
    lang = str(md.get("lang") or "").strip().lower()
    return {
        "milvus_pk": "" if row.get("id") is None else str(row.get("id")),
        "business_key": f"{recipe_id}:{lang}" if recipe_id and lang else "",
        "recipe_id": recipe_id,
        "lang": lang,
        "name": str(md.get("name") or "").strip(),
        "image_url": str(md.get("image_url") or "").strip(),
        "ingredients": _list_text(md.get("ingredients")),
        "ingredients_raw": _list_text(md.get("ingredients_raw")),
        "tags": _list_text(md.get("tags")),
        "facets": _json_text(md.get("facets")),
        "description": str(md.get("description") or "").strip(),
        "index_text": str(row.get("text") or "").strip(),
    }


def _style_title(cell):
    cell.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(vertical="center")


def _style_header(row):
    fill = PatternFill("solid", fgColor="D9EAF7")
    font = Font(name="Arial", bold=True, color="1F1F1F")
    border = Border(bottom=Side(style="thin", color="A6A6A6"))
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _set_widths(ws, widths: dict[str, float]):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _add_kv(ws, row: int, key: str, value):
    ws.cell(row=row, column=1, value=key)
    ws.cell(row=row, column=2, value=value)
    ws.cell(row=row, column=1).font = Font(name="Arial", bold=True)
    ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")


def _quality_issues(records: list[dict]) -> tuple[list[dict], list[str]]:
    keys = [r["business_key"] for r in records if r["business_key"]]
    key_counts = Counter(keys)
    duplicate_keys = [k for k, c in key_counts.items() if c > 1]
    dirty = []
    for idx, rec in enumerate(records, start=2):
        reasons = []
        if not rec["recipe_id"]:
            reasons.append("missing_recipe_id")
        if rec["lang"] not in {"zh", "en"}:
            reasons.append("invalid_lang")
        if not rec["name"]:
            reasons.append("missing_name")
        if not rec["ingredients"]:
            reasons.append("missing_ingredients")
        if not rec["image_url"]:
            reasons.append("missing_image_url")
        elif not _valid_url(rec["image_url"]):
            reasons.append("invalid_image_url")
        if rec["business_key"] in duplicate_keys:
            reasons.append("duplicate_business_key")
        if reasons:
            dirty.append({"excel_row": idx, "business_key": rec["business_key"], "reasons": reasons})
    return dirty, duplicate_keys


def export_xlsx(uri: str, collection: str, output: Path, batch_size: int) -> dict:
    client = build_milvus_client(uri)
    try:
        if collection not in client.list_collections():
            raise RuntimeError(f"collection not found: {collection} @ {uri}")
        client.load_collection(collection)
        stats = client.get_collection_stats(collection)
        records = [_normalize(row) for row in _iter_rows(client, collection, batch_size)]
    finally:
        try:
            client.close()
        except Exception:
            pass

    records.sort(key=lambda r: (r["lang"], r["recipe_id"], str(r["milvus_pk"])))
    dirty, duplicate_keys = _quality_issues(records)
    lang_counts = Counter(r["lang"] for r in records)
    key_count = len({r["business_key"] for r in records if r["business_key"]})

    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "导出说明"
    ws_data = wb.create_sheet("recipe_hybrid_export")

    ws_info.merge_cells("A1:F1")
    ws_info["A1"] = f"CookClaw Milvus {collection} 导出"
    _style_title(ws_info["A1"])
    ws_info.row_dimensions[1].height = 28

    summary = [
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Milvus URI", uri),
        ("集合", collection),
        ("集合统计", _json_text(stats)),
        ("导出说明", "导出可读字段和 metadata；未展开 dense/sparse 向量字段。"),
        ("业务主键", "business_key = recipe_id + ':' + lang"),
        ("导出记录数", len(records)),
        ("中文 zh 记录数", lang_counts.get("zh", 0)),
        ("英文 en 记录数", lang_counts.get("en", 0)),
        ("业务主键数", key_count),
        ("重复业务主键数", len(duplicate_keys)),
        ("质量问题记录数", len(dirty)),
    ]
    row = 3
    for key, value in summary:
        _add_kv(ws_info, row, key, value)
        row += 1

    row += 1
    ws_info.cell(row=row, column=1, value="字段名")
    ws_info.cell(row=row, column=2, value="来源")
    ws_info.cell(row=row, column=3, value="说明")
    _style_header(ws_info[row])
    schema = [
        ("milvus_pk", "Milvus id", "Milvus 自增主键，不是业务主键。"),
        ("business_key", "metadata.recipe_id + metadata.lang", "业务主键，用于校验同一语言下唯一。"),
        ("recipe_id", "metadata.recipe_id", "菜谱业务 ID。"),
        ("lang", "metadata.lang", "语言标记 zh/en。"),
        ("name", "metadata.name", "菜谱名称。"),
        ("image_url", "metadata.image_url", "菜谱图片地址。"),
        ("ingredients", "metadata.ingredients", "清洗后的食材数组，导出为中文顿号分隔文本。"),
        ("ingredients_raw", "metadata.ingredients_raw", "迁移前原始食材字段，导出为文本。"),
        ("tags", "metadata.tags", "标签列表，导出为中文顿号分隔文本。"),
        ("facets", "metadata.facets", "结构化过滤维度，JSON 文本。"),
        ("description", "metadata.description", "描述字段，当前多数为空。"),
        ("index_text", "text", "BM25 索引用文本，名称 + 食材 + 标签。"),
    ]
    for item in schema:
        row += 1
        for col, value in enumerate(item, 1):
            ws_info.cell(row=row, column=col, value=value)
            ws_info.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    _set_widths(ws_info, {"A": 22, "B": 38, "C": 62, "D": 12, "E": 12, "F": 12})
    ws_info.freeze_panes = "A3"
    ws_info.sheet_view.showGridLines = False

    ws_data.append(HEADERS)
    _style_header(ws_data[1])
    for rec in records:
        ws_data.append([rec[h] for h in HEADERS])

    table_ref = f"A1:{get_column_letter(len(HEADERS))}{len(records) + 1}"
    table = Table(displayName="recipe_hybrid_export_table", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws_data.add_table(table)
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = table_ref
    _set_widths(ws_data, {
        "A": 14, "B": 28, "C": 22, "D": 9, "E": 28, "F": 62,
        "G": 44, "H": 52, "I": 66, "J": 48, "K": 28, "L": 72,
    })
    for row_cells in ws_data.iter_rows(min_row=2, max_row=ws_data.max_row):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        image_cell = row_cells[5]
        if image_cell.value:
            image_cell.hyperlink = image_cell.value
            image_cell.style = "Hyperlink"
    ws_data.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {
        "output": str(output),
        "records": len(records),
        "stats": stats,
        "lang_counts": dict(lang_counts),
        "duplicate_business_keys": len(duplicate_keys),
        "dirty": len(dirty),
        "dirty_samples": dirty[:5],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="导出 Milvus 菜谱集合为 Excel")
    parser.add_argument("--uri", default=MILVUS_CONFIG["uri"], help="Milvus URI，默认读取 .env")
    parser.add_argument("--collection", default=os.getenv("MILVUS_HYBRID_COLLECTION", MILVUS_CONFIG["hybrid_collection"]))
    parser.add_argument("--out", default=str(ROOT / "outputs" / "milvus_recipe_hybrid_export_2026-07-03.xlsx"))
    parser.add_argument("--batch-size", type=int, default=1000)
    args = parser.parse_args()

    result = export_xlsx(args.uri, args.collection, Path(args.out), args.batch_size)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
